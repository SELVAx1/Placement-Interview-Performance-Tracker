import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from config import TEMPLATES_DIR

def style_excel_sheet(ws, title: str):
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

def generate_templates():
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    print("Generating bulk upload sample templates in:", TEMPLATES_DIR)

    # ==============================================================
    # 1. TEMPLATE: DRIVE SHORTLIST (SHORTLIST MODE - EMAILS ONLY)
    # ==============================================================
    shortlist_headers = ["Student Gmail", "Student Name", "Branch"]
    shortlist_data = [
        ["rahul.sharma@college.edu", "Rahul Sharma", "CSE"],
        ["ananya.reddy@college.edu", "Ananya Reddy", "CSE"],
        ["vikram.patel@college.edu", "Vikram Patel", "IT"],
        ["deepa.krishnan@college.edu", "Deepa Krishnan", "ECE"],
        ["sneha.gupta@college.edu", "Sneha Gupta", "CSE"],
        ["arjun.menon@college.edu", "Arjun Menon", "ECE"],
        ["priya.nair@college.edu", "Priya Nair", "CSE"],
    ]

    # Save Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shortlisted Candidates"
    ws.append(shortlist_headers)
    for row in shortlist_data:
        ws.append(row)
    style_excel_sheet(ws, "Shortlist")
    wb.save(os.path.join(TEMPLATES_DIR, "sample_drive_shortlist.xlsx"))

    # Save CSV
    with open(os.path.join(TEMPLATES_DIR, "sample_drive_shortlist.csv"), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(shortlist_headers)
        writer.writerows(shortlist_data)

    print("[OK] sample_drive_shortlist.xlsx & .csv generated")

    # ==============================================================
    # 2. TEMPLATE: DRIVE RESULTS / VERDICTS (VERDICT MODE)
    # ==============================================================
    results_headers = ["Student Gmail", "Student Name", "Round", "Score", "Result Status"]
    results_data = [
        ["rahul.sharma@college.edu", "Rahul Sharma", 2, 78.5, "Shortlisted for Round 3"],
        ["ananya.reddy@college.edu", "Ananya Reddy", 3, 92.0, "Selected"],
        ["vikram.patel@college.edu", "Vikram Patel", 1, 35.0, "Rejected"],
        ["deepa.krishnan@college.edu", "Deepa Krishnan", 2, 45.0, "Rejected"],
        ["sneha.gupta@college.edu", "Sneha Gupta", 4, 95.0, "Selected"],
        ["arjun.menon@college.edu", "Arjun Menon", 2, 65.0, "On Hold"],
        ["priya.nair@college.edu", "Priya Nair", 3, 84.0, "Selected"],
    ]

    # Save Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Drive Results"
    ws.append(results_headers)
    for row in results_data:
        ws.append(row)
    style_excel_sheet(ws, "Results")

    from openpyxl.worksheet.datavalidation import DataValidation
    dv_res = DataValidation(type="list", formula1='"Shortlisted for Round 2,Shortlisted for Round 3,Selected,Rejected,On Hold,Absent"', allow_blank=True)
    dv_res.error = 'Please select a valid result status.'
    dv_res.errorTitle = 'Invalid Status'
    ws.add_data_validation(dv_res)
    dv_res.add("E2:E100")

    wb.save(os.path.join(TEMPLATES_DIR, "sample_drive_results.xlsx"))

    # Save CSV
    with open(os.path.join(TEMPLATES_DIR, "sample_drive_results.csv"), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(results_headers)
        writer.writerows(results_data)

    print("[OK] sample_drive_results.xlsx & .csv generated")

    # ==============================================================
    # 3. TEMPLATE: USER ACCESS & ROLES
    # ==============================================================
    user_headers = ["User Email", "Role", "Password"]
    user_data = [
        ["student.demo1@college.edu", "Student", "Pass@123"],
        ["student.demo2@college.edu", "Student", ""],
        ["prof.arun@college.edu", "Mentor", "Mentor#456"],
        ["prof.meena@college.edu", "Mentor", ""],
        ["coord.placement@college.edu", "Coordinator", "Coord@2026"],
        ["recruiter.zoho@company.com", "Recruiter", "Recruiter#789"],
        ["dept.head.cse@college.edu", "Department", "DeptHead@123"]
    ]

    # Save Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "User Access"
    ws.append(user_headers)
    for row in user_data:
        ws.append(row)
    style_excel_sheet(ws, "Users")

    from openpyxl.worksheet.datavalidation import DataValidation
    dv_user = DataValidation(type="list", formula1='"Student,Mentor,Coordinator,Department,Recruiter"', allow_blank=True)
    dv_user.error = 'Please select a valid user role.'
    dv_user.errorTitle = 'Invalid Role'
    ws.add_data_validation(dv_user)
    dv_user.add("B2:B100")

    wb.save(os.path.join(TEMPLATES_DIR, "sample_user_access.xlsx"))

    # Save CSV
    with open(os.path.join(TEMPLATES_DIR, "sample_user_access.csv"), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(user_headers)
        writer.writerows(user_data)

    print("[OK] sample_user_access.xlsx & .csv generated")

    # ==============================================================
    # 4. TEMPLATE: STUDENT ACADEMIC PROFILES ROSTER
    # ==============================================================
    roster_headers = [
        "Register Number", "Full Name", "Student Email", "Department", 
        "CGPA", "10th Percentage", "12th Percentage", "Technical Skills"
    ]
    roster_data = [
        ["2021CS101", "Rahul Sharma", "rahul.sharma@college.edu", "CSE", 7.8, 89.5, 85.2, "Python, SQL, Java"],
        ["2021CS102", "Ananya Reddy", "ananya.reddy@college.edu", "CSE", 8.5, 92.0, 88.0, "Java, Spring Boot, React"],
        ["2021IT103", "Vikram Patel", "vikram.patel@college.edu", "IT", 6.9, 78.0, 72.0, "Python, HTML, CSS"],
        ["2021EC104", "Deepa Krishnan", "deepa.krishnan@college.edu", "ECE", 7.2, 85.0, 80.0, "C++, Embedded Systems"],
        ["2021CS105", "Sneha Gupta", "sneha.gupta@college.edu", "CSE", 9.1, 95.0, 93.5, "DSA, System Design, C++, AWS"],
        ["2021EC106", "Arjun Menon", "arjun.menon@college.edu", "ECE", 7.5, 88.0, 82.0, "C, Python, IoT"],
        ["2021CS107", "Priya Nair", "priya.nair@college.edu", "CSE", 8.2, 90.5, 87.0, "Python, ML, Data Science"]
    ]

    # Save Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Roster"
    ws.append(roster_headers)
    for row in roster_data:
        ws.append(row)
    style_excel_sheet(ws, "Roster")
    wb.save(os.path.join(TEMPLATES_DIR, "sample_student_roster.xlsx"))

    # Save CSV
    with open(os.path.join(TEMPLATES_DIR, "sample_student_roster.csv"), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(roster_headers)
        writer.writerows(roster_data)

    print("[OK] sample_student_roster.xlsx & .csv generated")

    # ==============================================================
    # 5. TEMPLATE: COMPANY PLACEMENT DRIVES
    # ==============================================================
    drives_headers = [
        "Company Name", "Job Role", "CTC LPA", "Company Type", 
        "Required CGPA", "Allowed Branches", "Total Rounds", 
        "Location", "Drive Date", "Status"
    ]
    drives_data = [
        ["Google", "Software Engineer", 24.5, "Product", 8.0, "CSE, IT, ECE", 4, "On Campus", "2026-10-20", "Active"],
        ["Zoho Corporation", "Member Technical Staff", 8.4, "Product", 7.0, "All", 5, "On Campus", "2026-10-25", "Active"],
        ["TCS", "Systems Engineer", 7.0, "Service", 6.5, "All", 4, "Virtual", "2026-10-30", "Active"],
        ["Microsoft", "Software Development Engineer", 28.0, "Product", 8.5, "CSE, IT", 4, "On Campus", "2026-11-05", "Active"],
        ["Accenture", "Advanced App Engineering Analyst", 6.5, "Consulting", 6.0, "All", 3, "On Campus", "2026-11-12", "Active"],
        ["Amazon", "Cloud Support Associate", 14.0, "Product", 7.5, "CSE, IT, ECE, EEE", 4, "Virtual", "2026-11-18", "Active"]
    ]

    # Save Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Company Drives"
    ws.append(drives_headers)
    for row in drives_data:
        ws.append(row)
    style_excel_sheet(ws, "Drives")

    from openpyxl.worksheet.datavalidation import DataValidation
    dv_ctype = DataValidation(type="list", formula1='"Product,Service,Consulting,Startup"', allow_blank=True)
    ws.add_data_validation(dv_ctype)
    dv_ctype.add("D2:D100")

    dv_status = DataValidation(type="list", formula1='"Active,Upcoming,Completed,Archived"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add("J2:J100")

    wb.save(os.path.join(TEMPLATES_DIR, "sample_company_drives.xlsx"))

    # Save CSV
    with open(os.path.join(TEMPLATES_DIR, "sample_company_drives.csv"), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(drives_headers)
        writer.writerows(drives_data)

    print("[OK] sample_company_drives.xlsx & .csv generated")

if __name__ == "__main__":
    generate_templates()

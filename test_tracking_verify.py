from fastapi.testclient import TestClient
from app import app
import openpyxl, io

client = TestClient(app)

def test_all():
    # 1. Total cohort
    res = client.get('/api/coordinator/students-tracking')
    assert res.status_code == 200
    data = res.json()
    print('Total students fetched:', len(data['students']))
    print('Summary Stats:', data['stats'])

    # 2. Year-wise tests
    for yr in ['4th Year', '3rd Year', '2nd Year', '1st Year']:
        r = client.get(f'/api/coordinator/students-tracking?year={yr}')
        assert r.status_code == 200
        st_list = r.json()['students']
        assert len(st_list) > 0
        assert all(yr in s['year'] for s in st_list)
        print(f'Year filter [{yr}]: {len(st_list)} students verified.')

    # 3. Dept-wise tests
    for dept in ['CSE', 'IT', 'ECE', 'EEE', 'MECH']:
        r = client.get(f'/api/coordinator/students-tracking?department={dept}')
        assert r.status_code == 200
        st_list = r.json()['students']
        assert len(st_list) > 0
        assert all(s['department'].upper() == dept for s in st_list)
        print(f'Dept filter [{dept}]: {len(st_list)} students verified.')

    # 4. Status tests
    r_placed = client.get('/api/coordinator/students-tracking?status=Placed')
    assert r_placed.status_code == 200
    placed_list = r_placed.json()['students']
    assert len(placed_list) > 0
    assert all(s['placement_status'] == 'Placed' for s in placed_list)
    print(f'Status filter [Placed]: {len(placed_list)} students verified.')

    # 5. Process and Interventions integrity
    vikram = next(s for s in data['students'] if 'vikram' in s['email'])
    print('\n[Vikram Patel 360 Inspection]')
    print('Placement Status:', vikram['placement_status'])
    print('Drives attended:', len(vikram['process_history']))
    for p in vikram['process_history']:
        print('  - ' + p['company_name'] + ': Round ' + str(p['round']) + ' (' + str(p['result']) + '), Score: ' + str(p['score']))
    print('Interventions:', len(vikram['interventions']))
    for iv in vikram['interventions']:
        print('  - [' + str(iv['priority']) + '|' + str(iv['status']) + '] ' + str(iv['title']))
        print('    Action tasks: ' + str(len(iv['actions'])))

    # 6. Excel cohort export verification
    exp = client.get('/api/coordinator/export/students-tracking?year=4th%20Year&department=CSE')
    assert exp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(exp.content))
    ws = wb.active
    print('\nCohort Excel Export sheet title:', ws.title)
    print('Excel headers:', [cell.value for cell in ws[1]])
    print('Excel row count:', ws.max_row)

    # 7. Individual Student Dossier Endpoint & Excel Export Verification
    print('\n[Testing Individual Student Dossier & Excel Download]')
    reg_no = vikram['register_number']
    ind_res = client.get(f'/api/coordinator/student/{reg_no}')
    assert ind_res.status_code == 200
    ind_data = ind_res.json()
    assert ind_data['success'] is True
    assert ind_data['student']['register_number'] == reg_no
    assert ind_data['student']['name'] == vikram['name']
    print(f"Fetched individual student details for {vikram['name']} ({reg_no}) successfully.")

    # Individual Student Excel Dossier (.xlsx)
    ind_exp = client.get(f'/api/coordinator/export/student/{reg_no}')
    assert ind_exp.status_code == 200
    assert 'spreadsheetml' in ind_exp.headers.get('content-type', '')
    assert 'attachment' in ind_exp.headers.get('content-disposition', '')
    
    ind_wb = openpyxl.load_workbook(io.BytesIO(ind_exp.content))
    expected_sheets = [
        'Profile & Overview',
        'Interview Pipeline',
        'Interventions & Diagnostics',
        'Remediation Action Roadmap',
        'Mentor Notes & Coordinator Log'
    ]
    for sheet_name in expected_sheets:
        assert sheet_name in ind_wb.sheetnames, f"Missing sheet: {sheet_name}"
    print(f"Verified all 5 dossier worksheets in Excel workbook: {ind_wb.sheetnames}")
    
    # Check Profile sheet values
    ws_prof = ind_wb['Profile & Overview']
    assert ws_prof['A1'].value == 'STUDENT PLACEMENT & INTERVENTION DOSSIER'
    
    # Check Interview Pipeline sheet
    ws_pipe = ind_wb['Interview Pipeline']
    assert ws_pipe.max_row >= 2
    print(f"Pipeline worksheet rows: {ws_pipe.max_row}")

    # Check Interventions sheet
    ws_iv_sheet = ind_wb['Interventions & Diagnostics']
    assert ws_iv_sheet.max_row >= 2
    print(f"Interventions worksheet rows: {ws_iv_sheet.max_row}")

    # Check Action Roadmap sheet
    ws_act_sheet = ind_wb['Remediation Action Roadmap']
    assert ws_act_sheet.max_row >= 2
    print(f"Remediation Action Roadmap rows: {ws_act_sheet.max_row}")

    # Check email identifier lookup
    email_exp = client.get(f"/api/coordinator/export/student/{vikram['email']}")
    assert email_exp.status_code == 200

    # Check 404 on nonexistent student
    bad_res = client.get('/api/coordinator/student/UNKNOWN999')
    assert bad_res.status_code == 404

    print('\nALL COHORT & INDIVIDUAL STUDENT DOSSIER VERIFICATIONS PASSED 100% SUCCESSFULLY!')

if __name__ == '__main__':
    test_all()


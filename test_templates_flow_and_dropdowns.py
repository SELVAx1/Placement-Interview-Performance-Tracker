import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from app import app
import db

client = TestClient(app)

def test_template_candidate_details_dropdown_and_round_flow():
    """
    Test user requirement:
    1. If student appeared/applied in Round 1, Round 1 template MUST contain their details.
    2. Round 1 template MUST contain Excel dropdown options on Result Status / Verdict (Shortlisted for Round 2, Selected, Rejected, etc.).
    3. When coordinator updates the record with 'Shortlisted for Round 2', student clears Round 1 and MOVES to appear in Round 2 (not selected in Round 2).
    4. Student MUST now appear in Round 2 update template with dropdown options.
    5. When updated in subsequent rounds, candidate moves forward stage by stage until final round selection.
    6. Rejected candidates do NOT appear in subsequent round templates.
    """
    db.init_db()

    # 1. Create a 3-round drive
    drive_payload = {
        "company_name": "TemplateFlow Technologies",
        "job_role": "Full Stack Engineer",
        "ctc_lpa": 16.5,
        "min_cgpa": 7.0,
        "allowed_branches": "CSE, IT, ECE",
        "location": "Bengaluru",
        "total_rounds": 3,
        "description": "3-stage test recruitment drive."
    }
    create_res = client.post("/api/drives", json=drive_payload)
    assert create_res.status_code == 201
    drive = create_res.json()["drive"]
    drive_id = drive["id"]

    student_pos = "candidate_pass@gmail.com"
    student_neg = "candidate_fail@gmail.com"

    # 2. Both candidates apply for Round 1
    apply_res1 = client.post("/api/student/apply", json={"drive_id": drive_id, "gmail": student_pos})
    assert apply_res1.status_code == 200
    apply_res2 = client.post("/api/student/apply", json={"drive_id": drive_id, "gmail": student_neg})
    assert apply_res2.status_code == 200

    # 3. Download Round 1 Update Template (.xlsx)
    r1_tpl_res = client.get(f"/api/drives/{drive_id}/export/round/1/template")
    assert r1_tpl_res.status_code == 200
    
    wb_r1 = openpyxl.load_workbook(io.BytesIO(r1_tpl_res.content))
    ws_r1 = wb_r1.active

    # Check DataValidation dropdown exists
    assert len(ws_r1.data_validations.dataValidation) > 0, "Round 1 template MUST have openpyxl DataValidation dropdown"
    dv_r1 = ws_r1.data_validations.dataValidation[0]
    assert "Shortlisted for Round 2" in dv_r1.formula1, f"Dropdown must include Shortlisted for Round 2, got {dv_r1.formula1}"
    assert "Rejected" in dv_r1.formula1

    # Check student details in Round 1 template rows
    rows_r1 = list(ws_r1.iter_rows(values_only=True))
    header_r1 = rows_r1[0]
    assert "Student Gmail" in header_r1
    assert "Result Status / Verdict" in header_r1

    emails_r1 = [row[0].strip().lower() for row in rows_r1[1:] if row[0]]
    assert student_pos.lower() in emails_r1, "Applied student MUST appear in Round 1 template"
    assert student_neg.lower() in emails_r1, "Applied student 2 MUST appear in Round 1 template"

    # 4. Coordinator updates Round 1:
    # student_pos -> "Shortlisted for Round 2"
    # student_neg -> "Rejected"
    wb_upload_r1 = openpyxl.Workbook()
    ws_up1 = wb_upload_r1.active
    ws_up1.append(["Student Gmail", "Student Name", "Department", "Current Round", "Result Status / Verdict", "Score", "Feedback", "Weakness Area"])
    ws_up1.append([student_pos, "Candidate Pass", "CSE", "Round 1", "Shortlisted for Round 2", 88.0, "Great fundamentals", "None"])
    ws_up1.append([student_neg, "Candidate Fail", "ECE", "Round 1", "Rejected", 40.0, "Needs improvement", "Aptitude"])

    buf1 = io.BytesIO()
    wb_upload_r1.save(buf1)
    buf1.seek(0)

    upload_res1 = client.post(
        f"/api/drives/{drive_id}/upload-results",
        files={"file": ("r1_eval.xlsx", buf1.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert upload_res1.status_code == 200

    # Verify process funnel:
    # Round 1: Cleared = 1 (student_pos), Rejected = 1 (student_neg)
    # Round 2: Appeared = 1 (student_pos), Cleared = 0! (student_pos NOT selected in Round 2 yet!)
    proc1 = client.get(f"/api/drives/{drive_id}/process").json()
    r1_proc = proc1["rounds"][0]
    r2_proc = proc1["rounds"][1]

    assert r1_proc["cleared_count"] == 1
    assert r1_proc["rejected_count"] == 1
    assert r2_proc["appeared_count"] == 1
    assert r2_proc["cleared_count"] == 0, "Candidate must NOT be selected in Round 2"
    assert proc1["summary"]["selected_count"] == 0

    # 5. Check Round 2 Template:
    # student_pos MUST be present
    # student_neg (rejected) MUST NOT be present!
    r2_tpl_res = client.get(f"/api/drives/{drive_id}/export/round/2/template")
    assert r2_tpl_res.status_code == 200
    wb_r2 = openpyxl.load_workbook(io.BytesIO(r2_tpl_res.content))
    ws_r2 = wb_r2.active

    # Check DataValidation dropdown exists for Round 2 (points to Round 3)
    assert len(ws_r2.data_validations.dataValidation) > 0
    dv_r2 = ws_r2.data_validations.dataValidation[0]
    assert "Shortlisted for Round 3" in dv_r2.formula1

    rows_r2 = list(ws_r2.iter_rows(values_only=True))
    emails_r2 = [row[0].strip().lower() for row in rows_r2[1:] if row[0]]
    assert student_pos.lower() in emails_r2, "Cleared candidate MUST appear in Round 2 template"
    assert student_neg.lower() not in emails_r2, "Rejected candidate MUST NOT appear in Round 2 template"

    # 6. Coordinator updates Round 2: student_pos -> "Shortlisted for Round 3"
    wb_upload_r2 = openpyxl.Workbook()
    ws_up2 = wb_upload_r2.active
    ws_up2.append(["Student Gmail", "Student Name", "Department", "Current Round", "Result Status / Verdict", "Score", "Feedback", "Weakness Area"])
    ws_up2.append([student_pos, "Candidate Pass", "CSE", "Round 2", "Shortlisted for Round 3", 92.0, "Excellent system design", "None"])

    buf2 = io.BytesIO()
    wb_upload_r2.save(buf2)
    buf2.seek(0)

    upload_res2 = client.post(
        f"/api/drives/{drive_id}/upload-results",
        files={"file": ("r2_eval.xlsx", buf2.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert upload_res2.status_code == 200

    proc2 = client.get(f"/api/drives/{drive_id}/process").json()
    assert proc2["rounds"][1]["cleared_count"] == 1
    assert proc2["rounds"][2]["appeared_count"] == 1
    assert proc2["rounds"][2]["cleared_count"] == 0, "Must NOT be selected in Round 3 yet"

    # 7. Check Round 3 Template:
    # Round 3 is the final round!
    r3_tpl_res = client.get(f"/api/drives/{drive_id}/export/round/3/template")
    assert r3_tpl_res.status_code == 200
    wb_r3 = openpyxl.load_workbook(io.BytesIO(r3_tpl_res.content))
    ws_r3 = wb_r3.active

    assert len(ws_r3.data_validations.dataValidation) > 0
    dv_r3 = ws_r3.data_validations.dataValidation[0]
    assert "Selected" in dv_r3.formula1

    # 8. Coordinator selects student in Round 3 (Final Round)
    wb_upload_r3 = openpyxl.Workbook()
    ws_up3 = wb_upload_r3.active
    ws_up3.append(["Student Gmail", "Student Name", "Department", "Current Round", "Result Status / Verdict", "Score", "Feedback", "Weakness Area"])
    ws_up3.append([student_pos, "Candidate Pass", "CSE", "Round 3", "Selected", 95.0, "Outstanding leadership fit", "None"])

    buf3 = io.BytesIO()
    wb_upload_r3.save(buf3)
    buf3.seek(0)

    upload_res3 = client.post(
        f"/api/drives/{drive_id}/upload-results",
        files={"file": ("r3_eval.xlsx", buf3.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert upload_res3.status_code == 200

    proc3 = client.get(f"/api/drives/{drive_id}/process").json()
    assert proc3["rounds"][2]["cleared_count"] == 1
    assert proc3["summary"]["selected_count"] == 1
    assert len(proc3["selected_students"]) == 1
    assert proc3["selected_students"][0]["gmail"].lower() == student_pos.lower()

    # 9. Verify Selected Students Excel Export
    sel_export = client.get(f"/api/drives/{drive_id}/export/selected")
    assert sel_export.status_code == 200
    wb_sel = openpyxl.load_workbook(io.BytesIO(sel_export.content))
    ws_sel = wb_sel.active
    sel_text = " ".join([str(c.value or "") for row in ws_sel.iter_rows() for c in row])
    assert student_pos.lower() in sel_text.lower()
    assert student_neg.lower() not in sel_text.lower()

    # Clean up test rows
    conn = db.get_db_connection()
    c = conn.cursor()
    c.execute("DELETE FROM student_drive_results WHERE LOWER(gmail) IN (?, ?)", (student_pos.lower(), student_neg.lower()))
    conn.commit()
    conn.close()

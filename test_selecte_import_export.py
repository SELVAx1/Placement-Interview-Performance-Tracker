import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from app import app
import db

client = TestClient(app)

def test_selecte_import_and_round_advancement():
    """
    Test user request:
    When user uploads an Excel file with 'selecte' in the result column:
    1. It must be considered as 'Selected'.
    2. Round 'Round 2' must be correctly parsed as round 2.
    3. The candidate must be in cleared candidates for Round 2.
    4. The candidate must be in Round 3 evaluation template.
    5. The candidate must be in Selected Students export.
    6. The candidate must be considered Placed in Coordinator tracking.
    """
    db.init_db()

    # 1. Fetch drives
    drives_res = client.get("/api/drives")
    assert drives_res.status_code == 200
    drives = drives_res.json()["drives"]
    assert len(drives) > 0
    target_drive = next((d for d in drives if d.get("total_rounds") == 2), None)
    if not target_drive:
        create_res = client.post("/api/drives", json={
            "company_name": "QuickHire TwoRound",
            "job_role": "Analyst",
            "ctc_lpa": 10.0,
            "total_rounds": 2
        })
        target_drive = create_res.json()["drive"]
    drive_id = target_drive["id"]

    test_email = "test_selecte_candidate@gmail.com"
    test_email_neg = "test_rejected_candidate@gmail.com"

    # Clean up test email from DB
    conn = db.get_db_connection()
    c = conn.cursor()
    c.execute("DELETE FROM student_drive_results WHERE LOWER(gmail) IN (?, ?)", (test_email.lower(), test_email_neg.lower()))
    conn.commit()
    conn.close()

    # 2. Build Excel update spreadsheet with "selecte" and "Round 2"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Round 2 Update"
    ws.append(["Student Gmail", "Student Name", "Department", "Current Round", "Result Status / Verdict", "Score", "Feedback", "Weakness Area"])
    ws.append([test_email, "Selecte Test Student", "CSE", "Round 2", "selecte", 92.5, "Strong analytical skills", "DSA trees"])
    ws.append([test_email_neg, "Rejected Test Student", "ECE", "Round 2", "not selected", 45.0, "Needs more practice", "OOP concepts"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # 3. Upload to /api/drives/{drive_id}/upload-results
    files = {
        "file": ("round_2_eval_update.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    upload_res = client.post(f"/api/drives/{drive_id}/upload-results", files=files)
    assert upload_res.status_code == 200
    up_data = upload_res.json()
    assert up_data["success"] is True
    assert up_data["updated_count"] == 2

    # 4. Check results in DB
    results_res = client.get(f"/api/drives/{drive_id}/results")
    assert results_res.status_code == 200
    all_results = results_res.json()["results"]

    candidate_rec = next((r for r in all_results if r["gmail"].lower() == test_email.lower()), None)
    neg_rec = next((r for r in all_results if r["gmail"].lower() == test_email_neg.lower()), None)

    assert candidate_rec is not None, "Candidate record was not saved"
    assert neg_rec is not None, "Negative candidate record was not saved"

    # 'selecte' normalized to 'Selected'
    assert candidate_rec["result"] == "Selected", f"Expected 'Selected', got {candidate_rec['result']}"
    # 'Round 2' parsed to 2
    assert candidate_rec["round"] == 2, f"Expected round 2, got {candidate_rec['round']}"
    assert candidate_rec["score"] == 92.5
    assert candidate_rec["feedback"] == "Strong analytical skills"

    # 'not selected' normalized to 'Rejected'
    assert neg_rec["result"] == "Rejected", f"Expected 'Rejected', got {neg_rec['result']}"

    # 5. Check Process Details funnel
    proc_res = client.get(f"/api/drives/{drive_id}/process")
    assert proc_res.status_code == 200
    p_data = proc_res.json()

    # Verify Round 2 cleared students includes test_email
    r2_info = next((r for r in p_data["rounds"] if r["round_number"] == 2), None)
    assert r2_info is not None
    cleared_r2_emails = [s["gmail"].lower() for s in r2_info["cleared_students"]]
    assert test_email.lower() in cleared_r2_emails, f"{test_email} should be cleared in Round 2"
    assert test_email_neg.lower() not in cleared_r2_emails, f"{test_email_neg} should NOT be cleared"

    # Verify Selected students includes test_email
    selected_emails = [s["gmail"].lower() for s in p_data["selected_students"]]
    assert test_email.lower() in selected_emails, f"{test_email} should be in selected_students"
    assert test_email_neg.lower() not in selected_emails

    # 6. Check Round 3 Template export: candidate MUST appear in Round 3 template
    r3_tpl_res = client.get(f"/api/drives/{drive_id}/export/round/3/template")
    assert r3_tpl_res.status_code == 200
    wb_r3 = openpyxl.load_workbook(io.BytesIO(r3_tpl_res.content))
    ws_r3 = wb_r3.active
    r3_emails = [str(cell.value or "").strip().lower() for cell in [row[0] for row in list(ws_r3.iter_rows())[1:]]]
    assert test_email.lower() in r3_emails, f"Candidate with 'selecte' in Round 2 should be in Round 3 template, found {r3_emails}"
    assert test_email_neg.lower() not in r3_emails, "Rejected candidate should not be in Round 3 template"

    # 7. Check Selected Students Excel export
    sel_export = client.get(f"/api/drives/{drive_id}/export/selected")
    assert sel_export.status_code == 200
    wb_sel = openpyxl.load_workbook(io.BytesIO(sel_export.content))
    ws_sel = wb_sel.active
    # In selected export, Student Gmail is in column index 4 (0-based)
    sel_rows_text = " ".join([str(cell.value or "") for row in ws_sel.iter_rows() for cell in row])
    assert test_email.lower() in sel_rows_text.lower(), f"Selected candidate should be in selected export file"

    # 8. Check Coordinator Student Tracking: must be marked Placed
    track_res = client.get(f"/api/coordinator/students-tracking")
    assert track_res.status_code == 200
    t_data = track_res.json()
    cand_track = next((s for s in t_data["students"] if s["email"].lower() == test_email.lower()), None)
    if cand_track:
        assert cand_track["placement_status"] == "Placed", f"Expected Placed, got {cand_track['placement_status']}"

    # Clean up test rows
    conn = db.get_db_connection()
    c = conn.cursor()
    c.execute("DELETE FROM student_drive_results WHERE LOWER(gmail) IN (?, ?)", (test_email.lower(), test_email_neg.lower()))
    conn.commit()
    conn.close()

if __name__ == "__main__":
    test_selecte_import_and_round_advancement()
    print("ALL 'selecte' import/export tests PASSED successfully!")

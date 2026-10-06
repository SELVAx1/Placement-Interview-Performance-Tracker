import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from app import app
import db

client = TestClient(app)

def test_round_selection_and_final_offer_flow():
    """
    Test user requirements:
    1. In the round update template, appeared candidates appear with status NOT selected (empty Column E).
    2. Column E contains DataValidation dropdown with Selected and Rejected.
    3. Conditional formatting highlights Selected in Green and Rejected in Red.
    4. When coordinator imports with 'Selected' in an intermediate round (e.g. Round 1 of 3):
       - Candidate clears Round 1 and moves to appear in Round 2.
       - In Round 2, candidate appears as 'appeared' but NOT selected in Round 2 yet (cleared_count=0).
       - In Round 2 template, candidate appears with status NOT selected (empty Column E).
    5. This continues for all intermediate rounds.
    6. When coordinator imports with 'Selected' in the FINAL round:
       - Candidate clears final round and is OFFERED (status 'Offered' / placed).
       - Candidate is in selected_students and coordinator tracking marks them Placed.
    7. Candidates marked 'Rejected' stay in that round and do not advance.
    """
    db.init_db()

    # 1. Create a 3-round drive
    drive_payload = {
        "company_name": "PakkaCorp Tech",
        "job_role": "Backend Software Engineer",
        "ctc_lpa": 18.0,
        "min_cgpa": 7.0,
        "allowed_branches": "CSE, IT, ECE",
        "location": "Hyderabad",
        "total_rounds": 3,
        "description": "3-round technical assessment and selection process."
    }
    create_res = client.post("/api/drives", json=drive_payload)
    assert create_res.status_code == 201
    drive = create_res.json()["drive"]
    drive_id = drive["id"]

    student_winner = "winner_candidate@gmail.com"
    student_rejected = "eliminated_candidate@gmail.com"

    # Clean up test rows
    conn = db.get_db_connection()
    c = conn.cursor()
    c.execute("DELETE FROM student_drive_results WHERE LOWER(gmail) IN (?, ?)", (student_winner.lower(), student_rejected.lower()))
    conn.commit()
    conn.close()

    # 2. Both candidates apply for Round 1
    apply_res1 = client.post("/api/student/apply", json={"drive_id": drive_id, "gmail": student_winner})
    assert apply_res1.status_code == 200
    apply_res2 = client.post("/api/student/apply", json={"drive_id": drive_id, "gmail": student_rejected})
    assert apply_res2.status_code == 200

    # 3. Export Round 1 Update Template
    r1_tpl_res = client.get(f"/api/drives/{drive_id}/export/round/1/template")
    assert r1_tpl_res.status_code == 200

    wb_r1 = openpyxl.load_workbook(io.BytesIO(r1_tpl_res.content))
    ws_r1 = wb_r1.active

    # Check DataValidation dropdown
    assert len(ws_r1.data_validations.dataValidation) > 0, "Template must have DataValidation dropdown"
    dv_r1 = ws_r1.data_validations.dataValidation[0]
    assert "Selected" in dv_r1.formula1
    assert "Rejected" in dv_r1.formula1

    # Check conditional formatting rules for Green Selected and Red Rejected
    cf_list = list(ws_r1.conditional_formatting)
    assert len(cf_list) >= 1, "Template must have conditional formatting applied"
    rules = cf_list[0].rules
    assert len(rules) >= 2, f"Must have at least 2 rules (Selected and Rejected), found {len(rules)}"
    rule_formulas = [str(r.formula[0]) for r in rules if r.formula]
    assert any("Selected" in f for f in rule_formulas), "Must have conditional formatting rule for Selected"
    assert any("Rejected" in f for f in rule_formulas), "Must have conditional formatting rule for Rejected"

    # Verify both students appear in Round 1 template with NOT SELECTED (empty Column E)
    rows_r1 = list(ws_r1.iter_rows(values_only=True))
    assert rows_r1[0][0] == "Student Gmail"
    assert rows_r1[0][4] == "Result Status / Verdict"

    r1_data = {row[0].strip().lower(): row[4] for row in rows_r1[1:] if row[0]}
    assert student_winner.lower() in r1_data, "student_winner must appear in Round 1 template"
    assert student_rejected.lower() in r1_data, "student_rejected must appear in Round 1 template"
    assert (r1_data[student_winner.lower()] or "") == "", "student_winner must appear NOT selected (empty Column E) initially"
    assert (r1_data[student_rejected.lower()] or "") == "", "student_rejected must appear NOT selected initially"

    # 4. Coordinator uploads Round 1:
    # student_winner -> 'Selected'
    # student_rejected -> 'Rejected'
    wb_up1 = openpyxl.Workbook()
    ws_up1 = wb_up1.active
    ws_up1.append(["Student Gmail", "Student Name", "Department", "Current Round", "Result Status / Verdict", "Score", "Feedback", "Weakness Area"])
    ws_up1.append([student_winner, "Winner Candidate", "CSE", "Round 1", "Selected", 85.0, "Great coding skills", "None"])
    ws_up1.append([student_rejected, "Eliminated Candidate", "ECE", "Round 1", "Rejected", 35.0, "Basic gaps", "DSA"])

    buf1 = io.BytesIO()
    wb_up1.save(buf1)
    buf1.seek(0)

    upload_res1 = client.post(
        f"/api/drives/{drive_id}/upload-results",
        files={"file": ("r1_eval.xlsx", buf1.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert upload_res1.status_code == 200

    # Check process funnel:
    # Round 1: Cleared = 1 (winner), Rejected = 1
    # Round 2: Appeared = 1 (winner), Cleared = 0 (winner NOT selected in Round 2 yet!)
    proc1 = client.get(f"/api/drives/{drive_id}/process").json()
    assert proc1["rounds"][0]["cleared_count"] == 1
    assert proc1["rounds"][0]["rejected_count"] == 1
    assert proc1["rounds"][1]["appeared_count"] == 1
    assert proc1["rounds"][1]["cleared_count"] == 0, "Candidate must NOT be selected in Round 2 yet"
    assert proc1["summary"]["selected_count"] == 0, "Candidate must NOT be offered after Round 1"

    # 5. Export Round 2 Update Template:
    # student_winner MUST appear with NOT SELECTED (empty Column E)
    # student_rejected MUST NOT appear
    r2_tpl_res = client.get(f"/api/drives/{drive_id}/export/round/2/template")
    assert r2_tpl_res.status_code == 200

    wb_r2 = openpyxl.load_workbook(io.BytesIO(r2_tpl_res.content))
    ws_r2 = wb_r2.active
    rows_r2 = list(ws_r2.iter_rows(values_only=True))
    r2_data = {row[0].strip().lower(): row[4] for row in rows_r2[1:] if row[0]}

    assert student_winner.lower() in r2_data, "Winner must appear in Round 2 template"
    assert student_rejected.lower() not in r2_data, "Rejected student must NOT appear in Round 2 template"
    assert (r2_data[student_winner.lower()] or "") == "", "Winner must appear NOT selected (empty Column E) in Round 2 template"

    # 6. Coordinator uploads Round 2: student_winner -> 'Selected'
    wb_up2 = openpyxl.Workbook()
    ws_up2 = wb_up2.active
    ws_up2.append(["Student Gmail", "Student Name", "Department", "Current Round", "Result Status / Verdict", "Score", "Feedback", "Weakness Area"])
    ws_up2.append([student_winner, "Winner Candidate", "CSE", "Round 2", "Selected", 90.0, "Great system design", "None"])

    buf2 = io.BytesIO()
    wb_up2.save(buf2)
    buf2.seek(0)

    upload_res2 = client.post(
        f"/api/drives/{drive_id}/upload-results",
        files={"file": ("r2_eval.xlsx", buf2.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert upload_res2.status_code == 200

    # Verify process funnel:
    # Round 2: Cleared = 1
    # Round 3: Appeared = 1, Cleared = 0 (NOT selected in Round 3 yet!)
    # Summary: selected_count = 0 (NOT offered yet!)
    proc2 = client.get(f"/api/drives/{drive_id}/process").json()
    assert proc2["rounds"][1]["cleared_count"] == 1
    assert proc2["rounds"][2]["appeared_count"] == 1
    assert proc2["rounds"][2]["cleared_count"] == 0
    assert proc2["summary"]["selected_count"] == 0

    # 7. Export Round 3 (Final Round) Update Template:
    # student_winner MUST appear with NOT SELECTED (empty Column E)
    r3_tpl_res = client.get(f"/api/drives/{drive_id}/export/round/3/template")
    assert r3_tpl_res.status_code == 200

    wb_r3 = openpyxl.load_workbook(io.BytesIO(r3_tpl_res.content))
    ws_r3 = wb_r3.active
    rows_r3 = list(ws_r3.iter_rows(values_only=True))
    r3_data = {row[0].strip().lower(): row[4] for row in rows_r3[1:] if row[0]}

    assert student_winner.lower() in r3_data, "Winner must appear in Round 3 template"
    assert (r3_data[student_winner.lower()] or "") == "", "Winner must appear NOT selected in Round 3 template"

    # 8. Coordinator uploads Round 3 (FINAL ROUND) with 'Selected':
    # Because Round 3 is the FINAL round, the student MUST BE OFFERED!
    wb_up3 = openpyxl.Workbook()
    ws_up3 = wb_up3.active
    ws_up3.append(["Student Gmail", "Student Name", "Department", "Current Round", "Result Status / Verdict", "Score", "Feedback", "Weakness Area"])
    ws_up3.append([student_winner, "Winner Candidate", "CSE", "Round 3", "Selected", 95.0, "Outstanding executive presentation", "None"])

    buf3 = io.BytesIO()
    wb_up3.save(buf3)
    buf3.seek(0)

    upload_res3 = client.post(
        f"/api/drives/{drive_id}/upload-results",
        files={"file": ("r3_eval.xlsx", buf3.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert upload_res3.status_code == 200

    # 9. Verify Final Round Offer:
    # Round 3: Cleared = 1!
    # selected_count = 1!
    # selected_students contains winner!
    proc3 = client.get(f"/api/drives/{drive_id}/process").json()
    assert proc3["rounds"][2]["cleared_count"] == 1
    assert proc3["summary"]["selected_count"] == 1
    assert len(proc3["selected_students"]) == 1
    assert proc3["selected_students"][0]["gmail"].lower() == student_winner.lower()

    # Verify status in database is 'Offered'
    results_res = client.get(f"/api/drives/{drive_id}/results").json()
    winner_res = next((r for r in results_res["results"] if r["gmail"].lower() == student_winner.lower()), None)
    assert winner_res is not None
    assert winner_res["round"] == 3
    assert winner_res["result"] in ["Offered", "Selected"]

    # Verify Coordinator tracking reports Placed
    track_res = client.get("/api/coordinator/students-tracking").json()
    winner_track = next((s for s in track_res["students"] if s["email"].lower() == student_winner.lower()), None)
    if winner_track:
        assert winner_track["placement_status"] == "Placed"

    # 10. Verify Selected Students Excel Export contains winner
    sel_export = client.get(f"/api/drives/{drive_id}/export/selected")
    assert sel_export.status_code == 200
    wb_sel = openpyxl.load_workbook(io.BytesIO(sel_export.content))
    ws_sel = wb_sel.active
    sel_text = " ".join([str(c.value or "") for row in ws_sel.iter_rows() for c in row])
    assert student_winner.lower() in sel_text.lower()
    assert student_rejected.lower() not in sel_text.lower()

    # Clean up test rows
    conn = db.get_db_connection()
    c = conn.cursor()
    c.execute("DELETE FROM student_drive_results WHERE LOWER(gmail) IN (?, ?)", (student_winner.lower(), student_rejected.lower()))
    conn.commit()
    conn.close()

if __name__ == "__main__":
    test_round_selection_and_final_offer_flow()
    print("ALL Round Selection and Final Offer Flow tests passed successfully!")
